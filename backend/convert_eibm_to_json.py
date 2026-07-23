# -*- coding: utf-8 -*-
import os
import re
import json
import docx
import django

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "core.settings")
django.setup()

from django.conf import settings
from api.models import Certificate, Category

docx_path = r"D:\Dənizçilik_İmtahanları\Ləyihənin İmtahan sorulari\EIBM_Test (1).docx"
doc = docx.Document(docx_path)

# Extract Answer Key first
answer_key = {}

in_answer_key = False
for p in doc.paragraphs:
    text = p.text.strip()
    if "Cavab Açarı" in text or "Cavab acari" in text or "Cavab Acari" in text:
        in_answer_key = True
        continue
    
    if in_answer_key and text:
        # e.g., "1. C 2. C 3. C 4. C 5. B"
        matches = re.findall(r'(\d+)[\.\)]\s*([A-D])', text)
        for q_num, ans in matches:
            answer_key[int(q_num)] = ans

print(f"Parsed {len(answer_key)} entries from Cavab Açarı.")

# Extract Questions & Options
questions = []
current_q = None
in_answer_key = False

def is_red(run):
    if run.font and run.font.color and run.font.color.rgb:
        color_str = str(run.font.color.rgb).upper()
        if color_str in ['FF0000', 'RED']:
            return True
    return False

for p in doc.paragraphs:
    text = p.text.strip()
    if "Cavab Açarı" in text or "Cavab acari" in text or "Cavab Acari" in text:
        in_answer_key = True
        break
    
    if not text or in_answer_key:
        continue
    
    q_match = re.match(r'^(\d+)[\.\)]\s*(.+)', text)
    opt_match = re.match(r'^([A-D])[\.\)]\s*(.+)', text)
    
    if q_match:
        q_num = int(q_match.group(1))
        q_title = q_match.group(2)
        current_q = {
            "id": f"q{q_num:03d}",
            "question_num": q_num,
            "question": q_title,
            "options": {},
            "correct_answer": None,
            "explanation": ""
        }
        questions.append(current_q)
    elif opt_match and current_q is not None:
        opt_key = opt_match.group(1)
        opt_val = opt_match.group(2)
        current_q["options"][opt_key] = opt_val
        
        # Check red color
        if any(is_red(r) for r in p.runs):
            current_q["correct_answer"] = opt_key

# Fill missing from Answer Key
for q in questions:
    q_num = q["question_num"]
    if not q["correct_answer"] and q_num in answer_key:
        q["correct_answer"] = answer_key[q_num]
    # Remove temporary field
    del q["question_num"]

print(f"Total questions parsed: {len(questions)}")
missing = [q for q in questions if not q["correct_answer"]]
print(f"Questions missing answer: {len(missing)}")

output_dir = settings.QUESTIONS_DIR / 'special'
output_dir.mkdir(parents=True, exist_ok=True)
json_file_rel = 'special/eibm.json'
json_path = output_dir / 'eibm.json'

output_data = {
    "certificate": "Əmniyyətli İdarəetmə Haqqında Beynəlxalq Məcəllə",
    "questions": questions
}

with open(json_path, 'w', encoding='utf-8') as f:
    json.dump(output_data, f, ensure_ascii=False, indent=2)

print(f"Saved to {json_path}")

# Update DB
category = Category.objects.get(name='Xüsusi hazırlıq şəhadətnamələri üzrə')
cert = Certificate.objects.filter(category=category, name__icontains='Əmniyyətli İdarəetmə').first()
if cert:
    cert.json_file = json_file_rel
    cert.save()
    print(f"Updated Certificate: ID {cert.id} ('{cert.name}') -> json_file='{json_file_rel}'")
else:
    print("Certificate not found in DB!")
