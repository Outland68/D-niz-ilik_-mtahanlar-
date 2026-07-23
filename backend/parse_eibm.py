# -*- coding: utf-8 -*-
import os
import docx

path1 = r"C:\Users\Omer\Downloads\EIBM_Test.docx"
path2 = r"C:\Users\Omer\Downloads\EIBM_Test (1).docx"

for path in [path1, path2]:
    if os.path.exists(path):
        print(f"=== Reading {path} ===")
        doc = docx.Document(path)
        for i, p in enumerate(doc.paragraphs[:30]):
            if p.text.strip():
                print(f"L{i}: {p.text}")
        print("Total paragraphs:", len(doc.paragraphs))
        print("Total tables:", len(doc.tables))
