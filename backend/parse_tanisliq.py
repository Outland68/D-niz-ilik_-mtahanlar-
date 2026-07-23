# -*- coding: utf-8 -*-
import docx

docx_path = r"D:\Dənizçilik_İmtahanları\Ləyihənin İmtahan sorulari\DƏNİZÇİLƏR ÜÇÜN TƏHLÜKƏSİZLİK ÜZRƏ TANIŞLIQ VƏ İLKİN HAZIRLIQ TESTİ.docx"
doc = docx.Document(docx_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")

for i, p in enumerate(doc.paragraphs[:40]):
    runs_info = []
    for r in p.runs:
        runs_info.append(f"['{r.text}', bold={r.bold}]")
    print(f"P{i+1}: {p.text.strip()} | Runs: {', '.join(runs_info)}")
