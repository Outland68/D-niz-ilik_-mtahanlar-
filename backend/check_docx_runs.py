import docx, os, sys

sys.stdout.reconfigure(encoding='utf-8')

src_dir = r"D:\Dənizçilik_İmtahanları\Ləyihənin İmtahan sorulari"

docx_files = [
    "DƏNİZÇİLƏR ÜÇÜN TƏHLÜKƏSİZLİK ÜZRƏ TANIŞLIQ VƏ İLKİN HAZIRLIQ TESTİ.docx",
    "EIBM_Test (1).docx",
    "Safety Familiarization & Basic Training Test.docx"
]

for df in docx_files:
    p = os.path.join(src_dir, df)
    doc = docx.Document(p)
    print(f"\n==================================================")
    print(f"DOCX: {df}")
    
    yellow_count = 0
    bold_count = 0
    red_count = 0
    
    for para in doc.paragraphs:
        for run in para.runs:
            if run.font.highlight_color:
                yellow_count += 1
            if run.bold:
                bold_count += 1
            if run.font.color and run.font.color.rgb:
                red_count += 1
                
    print(f"Highlighted runs: {yellow_count}, Bold runs: {bold_count}, Colored runs: {red_count}")
    
    # print sample question with runs
    for p_idx, para in enumerate(doc.paragraphs[:20]):
        if para.text.strip():
            runs_info = []
            for r in para.runs:
                if r.text.strip():
                    flags = []
                    if r.bold: flags.append('BOLD')
                    if r.font.highlight_color: flags.append(f'HL:{r.font.highlight_color}')
                    if r.font.color and r.font.color.rgb: flags.append(f'RGB:{r.font.color.rgb}')
                    flag_str = ','.join(flags) if flags else 'NORM'
                    runs_info.append(f"[{flag_str}] {r.text.strip()}")
            print(f"P{p_idx}: {' | '.join(runs_info)}")
